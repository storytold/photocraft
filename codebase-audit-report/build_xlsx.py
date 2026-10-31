import re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
R=__import__("os").path.dirname(__import__("os").path.abspath(__file__))+"/"
lines=open(R+'README.md').read().split('\n')

def table_at(n):  # 1-based line of header
    i=n-1; rows=[]
    while i<len(lines) and lines[i].startswith('|'):
        rows.append(lines[i]); i+=1
    cells=[[c.strip() for c in r.strip().strip('|').split('|')] for r in rows]
    return cells[0], cells[2:]

def table_after(head):
    i=[k for k,l in enumerate(lines) if l.startswith('#') and head in l][0]
    while not lines[i].startswith('|'): i+=1
    return table_at(i+1)

RANGE=re.compile(r'^\$([\d,]+) - \$([\d,]+)$')
def conv(s):
    bold=s.startswith('**'); s=s.replace('**','').replace('`','').strip()
    t=s.replace(',','')
    try:
        if re.fullmatch(r'-?\d+',t): return int(t),'int',bold
        if re.fullmatch(r'-?\d+\.\d+',t): return float(t),'float',bold
        if re.fullmatch(r'-?\d+(\.\d+)?%',t): return float(t[:-1])/100,'pct',bold
        if re.fullmatch(r'\$-?\d+',t): return int(t[1:]),'usd',bold
    except ValueError: pass
    return s,'text',bold

def norm(hdr,rows):
    # split $a - $b range columns
    rows=[r+['']*(len(hdr)-len(r)) for r in rows]
    out_h=[];split=[]
    for j,h in enumerate(hdr):
        if any(RANGE.match(r[j].replace('**','')) for r in rows): out_h+= [h.replace('**','')+' (low)',h.replace('**','')+' (high)']; split.append(j)
        else: out_h.append(h.replace('**','').replace('`',''))
    out=[]
    for r in rows:
        nr=[]
        for j,c in enumerate(r):
            if j in split:
                m=RANGE.match(c.replace('**',''))
                if m: nr+=[(int(m[1].replace(',','')),'usd',False),(int(m[2].replace(',','')),'usd',False)]
                else: nr+=[conv(c),('', 'text',False)]
            else: nr.append(conv(c))
        out.append(nr)
    return out_h,out

F='Arial'
HDR=PatternFill('solid',fgColor='1F3864'); BAND=PatternFill('solid',fgColor='EEF3FA'); TOT=PatternFill('solid',fgColor='D9E2F3')
thin=Side(style='thin',color='B4C6E7'); BOT=Border(bottom=thin)
FMT={'int':'#,##0','float':'#,##0.00','pct':'0.0%','usd':'$#,##0'}
wb=Workbook(); wb.remove(wb.active)
stats={'formulas':0}

def sheet(name,title,blocks,note=None):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
    ws['A1']=title; ws['A1'].font=Font(name=F,size=16,bold=True,color='1F3864')
    r=2
    if note: ws.cell(r,1,note).font=Font(name=F,size=9,italic=True,color='595959'); r+=1
    widths={}
    first=True
    for sub,(hdr,rows) in blocks:
        hdr,rows=norm(hdr,rows); r+=1
        if sub: ws.cell(r,1,sub).font=Font(name=F,size=12,bold=True,color='2F5496'); r+=1
        for j,h in enumerate(hdr,1):
            c=ws.cell(r,j,h); c.font=Font(name=F,bold=True,color='FFFFFF'); c.fill=HDR
            c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            widths[j]=max(widths.get(j,8),min(max(len(h)+2,11),18))
        ws.row_dimensions[r].height=30
        hr=r
        if first and len(rows)>12: ws.freeze_panes=ws.cell(r+1,1)
        start=r+1; seg=start
        for i,row in enumerate(rows):
            r+=1; bold=row[0][2]
            for j,(v,k,b) in enumerate(row,1):
                c=ws.cell(r,j,v if v!='' else None)
                c.font=Font(name=F,bold=bold or b,size=10)
                if k in FMT: c.number_format=FMT[k]; c.alignment=Alignment(horizontal='right')
                else:
                    c.alignment=Alignment(wrap_text=len(str(v))>60,vertical='top')
                    widths[j]=max(widths.get(j,8),min(len(str(v))+2,60))
                if k in FMT: widths[j]=max(widths.get(j,8),len(f'{v:,}')+4 if k!='pct' else 9)
                c.border=BOT
                c.fill=TOT if bold else (BAND if i%2 else PatternFill())
            if bold:
                # replace additive totals with SUM formulas when they reproduce the value
                for j,(v,k,b) in enumerate(row,1):
                    if k in ('int','usd') and r>seg:
                        vals=[ws.cell(x,j).value for x in range(seg,r)]
                        if all(isinstance(x,(int,float)) or x is None for x in vals) and sum(x or 0 for x in vals)==v:
                            ws.cell(r,j).value=f'=SUM({L(j)}{seg}:{L(j)}{r-1})'; stats['formulas']+=1
                seg=r+1
        if len(rows)>12: ws.auto_filter.ref=f'A{hr}:{L(len(hdr))}{r}'
        first=False; r+=1
    for j,w in widths.items(): ws.column_dimensions[L(j)].width=w
    return ws

T=table_after
# Summary sheet
ws=wb.create_sheet('Summary'); ws.sheet_view.showGridLines=False
ws['A1']='PhotoCraft Codebase Cartography Audit: Key Numbers'; ws['A1'].font=Font(name=F,size=18,bold=True,color='1F3864')
ws['A2']='Commit 366002f, measured 2026-10-14. Each value links to the sheet it comes from.'; ws['A2'].font=Font(name=F,italic=True,size=9,color='595959')
kpis=[]  # filled after sheets exist
# scc language table from code block
i=[k for k,l in enumerate(lines) if l.startswith('Language ') and 'Files' in l][0]
lang=[]
for l in lines[i+2:]:
    if l.startswith('───'): continue
    if l.startswith('```'): break
    p=re.split(r'\s{2,}',l.strip())
    lang.append(p)
hdr=['Language','Files','Lines','Blanks','Comments','Code','Complexity']
lrows=[p for p in lang if p[0]!='Total']+[['**Total**']+['**%s**'%format(sum(int(p[j].replace(',','')) for p in lang if p[0]!='Total'),',') for j in range(1,7)]]
sheet('Languages','Language breakdown (scc 3.5.0, whole repo)',[(None,(hdr,lrows))])
sheet('LOC Summary','Lines of code: System / DevOps-Test / Total',[(None,T('4.2 Lines of code'))])
sheet('Components','Per-component rollup (LOC, types, functions, cyclomatic complexity)',[(None,T('Per-component rollup'))],
      'Rust CC = extended McCabe; other languages = scc proxy. Subtotals are SUM formulas.')
sheet('LOC by Language','LOC by category and language (backfiring input)',[(None,T('LOC by category and language'))])
sheet('Complex Files','Top 15 most complex Rust files',[(None,T('Top 15 most complex'))])
cl=[l.split() for l in open(R+'metrics/clippy-by-lint.txt') if l.strip()]
sheet('Clippy','Clippy findings by lint (non-test code; pedantic + never-crash lints)',[(None,(['Lint','Count'],[[x[1],x[0]] for x in cl]+[['**Total**','**%d**'%sum(int(x[0]) for x in cl)]]))])
hl=lambda p:[k for k,l in enumerate(lines) if l.startswith(p)][0]+1
sheet('FP Transactions','Transaction functions (EI / EQ / EO) from the command registry',[('By type',table_at(hl('| Type | Count |'))),('By command prefix',table_at(hl('| Command prefix |')))])
sheet('Function Points','IFPUG function points (System)',[('Function types',T('IFPUG function-type table')),('General System Characteristics',T('General System Characteristics')),
   ('IFPUG vs backfiring',T('System reconciliation')),('Subtotals and total',T('Function points: subtotals'))])
sheet('Backfiring','Capers Jones backfiring (LOC to FP)',[('System',T('Backfiring - System')),('DevOps/Test',T('Backfiring - DevOps'))])
sheet('Cost','Economic valuation: effort, cost, schedule',[('Salary inputs',T('Salary inputs')),('Effort by productivity band',T('Effort by productivity band')),
   ('Cost by productivity band (fully loaded, burden 1.5x-1.75x)',T('Cost by productivity band')),('Schedule sanity check',T('Schedule sanity check'))],
   'Replacement cost, not market value. Salary figures are secondary citations; see README section 4.7.')
sheet('Test Pyramid','Test pyramid',[(None,T('Pyramid'))])
sheet('Test Suites','Test-suite matrix (cargo test -- --list, default features)',[(None,T('Test-suite matrix'))])
sheet('Tools','Tools and versions',[(None,T('1. How to reproduce'))])

# KPIs: find cells by label
def find(sh,label,col,last=False):
    hits=[row[0].row for row in wb[sh].iter_rows() if row[0].value is not None and str(row[0].value).startswith(label)]
    if not hits: sys.exit('missing '+label)
    return f"='{sh}'!{L(col)}{hits[-1] if last else hits[0]}" 
K=[('System LOC',find('LOC Summary','System',2),'#,##0'),('DevOps/Test LOC',find('LOC Summary','DevOps/Test',2),'#,##0'),
   ('Total authored LOC',find('LOC Summary','Total',2),'#,##0'),('System cyclomatic complexity',find('LOC Summary','System',7),'#,##0'),
   ('System function points (IFPUG AFP)',find('Function Points','IFPUG (AFP)',2),'#,##0'),('Backfired System FP',find('Function Points','Backfiring (point)',2),'#,##0'),
   ('Total FP (System + DevOps/Test)',find('Function Points','Total',3),'#,##0'),
   ('Headline cost, low (0.32 FP/day)',find('Cost','0.32',2,True),'$#,##0'),('Headline cost, high (0.32 FP/day)',find('Cost','0.32',3,True),'$#,##0'),
   ('Tests',find('Test Pyramid','Total',3),'#,##0'),('Clippy warnings',find('Clippy','Total',2),'#,##0')]
ws['A4']='Metric'; ws['B4']='Value'
for c in ('A4','B4'): ws[c].font=Font(name=F,bold=True,color='FFFFFF'); ws[c].fill=HDR
for n,(k,f,fmt) in enumerate(K,5):
    ws.cell(n,1,k).font=Font(name=F,size=11); c=ws.cell(n,2,f); c.number_format=fmt; c.font=Font(name=F,size=11,bold=True,color='008000')
    for x in (1,2): ws.cell(n,x).border=BOT; ws.cell(n,x).fill=BAND if n%2 else PatternFill()
ws.column_dimensions['A'].width=40; ws.column_dimensions['B'].width=18
from openpyxl.drawing.image import Image as XI
import glob
for p in sorted(glob.glob(R+'images/*.png'),key=lambda p:(p[-20:].count('annotated'),p)):
    nm=p.split('/')[-1][3:-4].replace('-',' ').title()
    d=wb.create_sheet(('Diagram '+nm)[:31]); d.sheet_view.showGridLines=False
    d['A1']=nm+' view'; d['A1'].font=Font(name=F,size=16,bold=True,color='1F3864')
    d['A2']='Source: codebase-audit-report/'+p.split('report/')[1]; d['A2'].font=Font(name=F,size=9,italic=True,color='595959')
    im=XI(p)
    if im.width>1400: im.height=int(im.height*1400/im.width); im.width=1400
    d.add_image(im,'A4')
for s in wb.worksheets: s.sheet_properties.tabColor='2F5496' if s.title!='Summary' else 'C00000'
out=R+'codebase-audit-metrics.xlsx'; wb.save(out); print(out,stats)
